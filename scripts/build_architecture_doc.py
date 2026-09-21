from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "Documento_Tecnico_Arquitectura_JARVIS.docx"
ASSET = ROOT / "tmp" / "architecture_document_assets"
NAVY, PALE, GRID, TEXT, MUTED = "183B63", "F4F7FA", "D9E2F3", "1F2933", "52616B"


def run(p, value, bold=False, size=10.3, color=TEXT, italic=False):
    r = p.add_run(value)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def spacing(p, before=0, after=6, line=1.12):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tc_pr.append(node)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    nodes = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRID)
        nodes.append(node)
    tc_pr.append(nodes)


def margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:tcMar")
    for edge, value in (("top", "85"), ("start", "115"), ("bottom", "85"), ("end", "115")):
        m = OxmlElement(f"w:{edge}")
        m.set(qn("w:w"), value)
        m.set(qn("w:type"), "dxa")
        node.append(m)
    tc_pr.append(node)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, value in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, NAVY); borders(c); margins(c); c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths: c.width = Inches(widths[i])
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, after=0)
        run(p, value, True, 8.7, "FFFFFF")
    head = OxmlElement("w:tblHeader"); head.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(head)
    for index, values in enumerate(rows):
        row = t.add_row()
        no_split = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(no_split)
        for i, value in enumerate(values):
            c = row.cells[i]
            shade(c, "FFFFFF" if index % 2 == 0 else PALE); borders(c); margins(c)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths: c.width = Inches(widths[i])
            p = c.paragraphs[0]; spacing(p, after=0, line=1.02)
            run(p, value, False, 8.35)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def heading(doc, value, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    spacing(p, before=15 if level == 1 else 9, after=5)
    run(p, value, True, 15 if level == 1 else 11.7, "000000")


def body(doc, value):
    p = doc.add_paragraph(); spacing(p, after=7); run(p, value)


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet"); spacing(p, after=3); run(p, value)


def caption(doc, value):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, after=9)
    run(p, value, False, 8.8, MUTED, True)


def page_number(p):
    run(p, "Página ", False, 8.5, MUTED)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    p.runs[-1]._r.extend([begin, instr, end])


def picture(doc, name):
    doc.add_picture(str(ASSET / name), width=Inches(6.8))


def set_base_fonts(doc):
    for style_name in ("Normal", "List Bullet", "Title", "Subtitle"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_base_fonts(doc)
    sec = doc.sections[0]
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2); sec.right_margin = Cm(2)
    sec.different_first_page_header_footer = True
    h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(h, "JARVIS | Documento Técnico de Arquitectura | v0.2.0", False, 8.5, MUTED)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER; page_number(f)

    # Portada académica solicitada. La cabecera y la paginación se omiten en esta primera página.
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, before=72, after=68)
    run(p, "Asistente Virtual: Jarvis", True, 24, "000000")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, after=5)
    run(p, "Integrantes:", True, 14, "000000")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, after=40)
    run(p, "Juan Esteban Rubio", False, 14, "000000")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, after=5)
    run(p, "Profesor:", True, 14, "000000")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, after=42)
    run(p, "Santiago Rivadeneira", False, 14, "000000")
    for text in ("Facultad de Ingeniería", "Desarrollo Móvil"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, after=5); run(p, text, False, 14, "000000")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; spacing(p, before=48, after=0)
    run(p, "Chía, 2026", False, 14, "000000")
    doc.add_page_break()

    heading(doc, "1 Propósito y alcance")
    body(doc, "Este documento describe la arquitectura implementada de JARVIS, un asistente móvil para organización académica y finanzas personales. La solución integra React Native y Expo en el teléfono, FastAPI en un servidor local, un orquestador con function calling, agentes especializados, PostgreSQL y Tailscale para conectividad privada.")
    body(doc, "El teléfono conserva la captura de voz, la síntesis de voz, la presentación de resultados y la confirmación explícita. El servidor concentra la lógica de negocio, la persistencia personal y la inferencia del modo evaluado. Esta separación evita que el móvil tenga acceso SQL o credenciales de base de datos.")
    heading(doc, "1.1 Alcance implementado", 2)
    for item in (
        "Aplicación React Native y Expo con dictado Android, TTS, respuestas estructuradas y estados de conversación.",
        "Backend REST FastAPI, autenticación Bearer, orquestador, agentes de Secretaría y Finanzas y propuestas idempotentes.",
        "PostgreSQL para perfil, recuerdos, conversaciones, agenda, tareas, transacciones, cuentas, pasivos, presupuestos y metas.",
        "Tailscale Serve como canal HTTPS privado y webhook bancario con token independiente.",
    ): bullet(doc, item)
    heading(doc, "1.2 Dependencias operativas", 2)
    for item in (
        "La rutina Android de cero fricción requiere una configuración real de Tasker, MacroDroid o NotificationListener con permisos concedidos.",
        "La lectura de correo requiere una cuenta IMAP configurada. El envío no es automático: se generan borradores sujetos a confirmación.",
        "Reconocimiento de voz, TTS y conectividad de la tailnet dependen del dispositivo y de sus servicios habilitados.",
    ): bullet(doc, item)

    heading(doc, "2 Trazabilidad con la rúbrica")
    table(doc, ["Criterio", "Arquitectura implementada", "Evidencia"], [
        ["Voz y aplicación móvil", "Expo, React Native, RecognizerIntent Android, expo-speech, interfaz de propuestas.", "Servicios de voz, conversación y pruebas de interfaz."],
        ["Orquestador multi agente", "FastAPI, catálogo cerrado de herramientas, Ollama function calling, rutas secretary, financial y composite.", "backend/app/orchestrator.py y agents.py."],
        ["Pagos cero fricción", "Webhook tipado, token dedicado, event_id idempotente y extracción JSON.", "bank_ingestion.py y contrato HTTP."],
        ["Tailscale y seguridad", "FastAPI en localhost, Tailscale Serve HTTPS, sin Funnel ni puertos públicos.", "Scripts de configuración y cliente móvil."],
        ["Persistencia relacional", "PostgreSQL con migraciones, FK, checks, índices y repositorios.", "database.py y repositories.py."],
    ], [1.3, 3.3, 1.9])
    doc.add_page_break()

    heading(doc, "3 Topología y fronteras de confianza")
    body(doc, "El móvil no accede directamente a la base de datos. Tailscale publica una ruta HTTPS exclusivamente dentro de la tailnet y hace proxy al backend que escucha en localhost. El webhook bancario usa la misma red privada, pero posee una credencial distinta de la aplicación móvil.")
    picture(doc, "topology.png")
    caption(doc, "Figura 1. Topología lógica, actores externos y fronteras de confianza.")
    table(doc, ["Capa", "Responsabilidad", "Restricción"], [
        ["Móvil", "Voz, UI, confirmación y preferencias.", "No ejecuta SQL remoto ni conserva secretos de base de datos."],
        ["Tailscale", "Canal mesh privado y proxy HTTPS.", "No habilita Funnel ni abre puertos del router."],
        ["FastAPI", "Validación HTTP, autenticación, propuestas y webhook.", "No permite que el LLM elija código o SQL arbitrario."],
        ["Orquestador y agentes", "Selección de herramientas y reglas de Secretaría o Finanzas.", "Las escrituras conversacionales requieren confirmación."],
        ["PostgreSQL", "Fuente de verdad del modo Servidor.", "No es accesible desde el teléfono."],
    ], [1.2, 3.15, 2.15])

    heading(doc, "4 Arquitectura del cliente móvil")
    body(doc, "Expo Router organiza Inicio, Tareas, JARVIS, Agenda y Perfil. Los componentes se concentran en presentación; engines y lib alojan cálculos puros; los servicios de almacenamiento encapsulan SQLite local, AsyncStorage y la adaptación del workspace remoto. WorkspaceProvider sustituye las vistas académicas y financieras por el snapshot autenticado cuando el modo Servidor está activo.")
    table(doc, ["Módulo", "Función", "Dependencia"], [
        ["src/app y src/features", "Pantallas, rutas, composición visual y comportamiento de presentación.", "React Native y Expo Router."],
        ["src/services/voice", "Dictado Android y respuesta hablada.", "Servicio de reconocimiento y TTS del sistema."],
        ["src/services/backend", "Cliente HTTP tipado, timeout, Bearer y errores comprensibles.", "FastAPI por URL Tailscale."],
        ["src/services/storage", "SQLite local, preferencias y caché mínima de identidad confirmada.", "Almacenamiento del dispositivo."],
        ["src/features/jarvis", "Mensajes, evidencias de herramientas, propuestas y reintentos.", "request_id y conversation_id."],
    ], [1.3, 3.25, 1.95])
    body(doc, "El deep link jarvis://talk abre la conversación y solicita dictado en una instalación nativa. Es un acceso rápido de voz y no reemplaza la automatización bancaria de cero fricción.")
    doc.add_page_break()

    heading(doc, "5 Secuencia de comando por voz")
    body(doc, "Los mensajes escritos y dictados comparten el mismo flujo. El cliente envía texto junto con request_id y conversation_id. El backend devuelve mensaje, ruta, resultados de herramientas y propuestas pendientes. Un cambio se persiste únicamente después de que el usuario confirma la propuesta.")
    picture(doc, "voice_sequence.png")
    caption(doc, "Figura 2. Secuencia de una orden por voz con confirmación y recarga de workspace.")
    heading(doc, "5.1 Consistencia e idempotencia", 2)
    for item in (
        "El request_id se reutiliza en un reintento para no ejecutar dos veces la misma operación.",
        "Las propuestas pasan por los estados pending, confirmed o cancelled y devuelven replayed cuando la confirmación ya fue aplicada.",
        "Tras confirmar, el móvil solicita el workspace del mes para reflejar el estado persistido del servidor.",
        "El orquestador no acepta que el modelo simule una propuesta en texto: debe invocar una herramienta para producir una confirmación válida.",
    ): bullet(doc, item)

    heading(doc, "6 Orquestador y agentes")
    body(doc, "El orquestador entrega al LLM un catálogo cerrado de funciones con esquemas Pydantic. Las llamadas se validan y se delegan a un agente. La ruta es Secretaría, Finanzas o Composite según las herramientas utilizadas. Los resultados externos se devuelven al LLM como datos delimitados, nunca como instrucciones ejecutables.")
    table(doc, ["Agente", "Capacidades", "Control"], [
        ["Secretaría", "Materias, tareas, agenda, recordatorios, lectura IMAP y borradores.", "Correo de solo lectura; borradores y cambios con confirmación."],
        ["Finanzas", "Ingresos, gastos, presupuesto, flujo de caja, tarjetas, préstamos e inversión de metas.", "Importes exactos, categorías y tipos validados."],
        ["Composite", "Combina respuestas académicas y financieras.", "Conserva las validaciones de cada dominio."],
    ], [1.15, 3.3, 2.05])

    heading(doc, "7 Ingestión bancaria de cero fricción")
    body(doc, "Tasker, MacroDroid o un NotificationListener puede detectar una notificación o correo bancario y enviar event_id, source y raw_text al webhook privado. El backend reserva el evento, solicita una extracción JSON de monto, moneda, comercio, fecha, medio de pago y categoría, y registra la transacción solo si la extracción satisface el umbral y las validaciones.")
    table(doc, ["Paso", "Contrato", "Resultado"], [
        ["Detección", "Regla Android para notificación o correo de la entidad.", "Texto disponible sin guardar credenciales bancarias."],
        ["Envío", "POST /v1/webhooks/bank-events, X-Jarvis-Webhook-Token y event_id.", "Canal HTTPS privado por Tailscale."],
        ["Extracción", "Salida JSON estructurada con confidence y campos tipados.", "Objeto validado o evento ignored."],
        ["Persistencia", "Reserva idempotente y vínculo único evento-transacción.", "created, duplicate o ignored; sin duplicados."],
    ], [1.1, 3.7, 1.7])
    body(doc, "El webhook constituye autorización previa de la automatización autenticada y puede insertar directamente. Esta excepción no aplica a órdenes conversacionales, que siempre se presentan para confirmación humana.")
    doc.add_page_break()

    heading(doc, "8 Despliegue y conectividad privada")
    body(doc, "FastAPI escucha por defecto en 127.0.0.1:8787. Tailscale Serve expone una URL HTTPS privada y reenvía a ese puerto local. El Android debe tener Tailscale activo y estar en la misma tailnet para resolver MagicDNS. Esta configuración permite operar desde datos móviles sin exponer el backend a Internet público.")
    table(doc, ["Elemento", "Configuración", "Verificación"], [
        ["Backend", "npm run backend:start con secretos DPAPI y conexión PostgreSQL.", "GET /health retorna status ok."],
        ["Tailscale", "npm run backend:tailscale configura Serve hacia 8787.", "tailscale serve status indica tailnet only."],
        ["Móvil", "Modo Servidor con URL MagicDNS y token enmascarado.", "GET /v1/agents/status desde el teléfono."],
        ["Ollama", "URL y modelo configurables detrás del backend.", "Estado available, unavailable o error."],
    ], [1.1, 3.65, 1.75])
    heading(doc, "8.1 Manejo de fallos", 2)
    for item in (
        "El cliente diferencia timeout, error de red y error de resolución de host privado; este último indica que Tailscale debe activarse en el teléfono.",
        "La app conserva solo nombre, id y zona horaria confirmados para mantener el saludo sin conexión. Memorias, tareas y finanzas siguen en el servidor.",
        "El backend no inicia sin tokens y base de datos válidos; los argumentos de las herramientas se validan antes de persistir.",
    ): bullet(doc, item)

    heading(doc, "9 Modelo relacional implementado")
    body(doc, "PostgreSQL es la fuente de verdad en modo Servidor. SQLAlchemy y los repositorios encapsulan acceso; las migraciones son transaccionales y versionadas. SQLite permanece en el móvil para modo local y pruebas aisladas, no como réplica de los datos personales del servidor.")
    picture(doc, "relational_model.png")
    caption(doc, "Figura 3. Entidades principales y relaciones del modelo de datos implementado.")
    table(doc, ["Grupo", "Relaciones", "Integridad"], [
        ["Identidad y memoria", "users 1 a 1 user_profiles; users 1 a N memories y conversations.", "FK, límites de longitud, estados, importancia y borrado lógico."],
        ["Académico", "academic_subjects 1 a N tasks y calendar_events.", "Nombre normalizado único, prioridad y estado restringidos."],
        ["Finanzas", "financial_accounts 1 a N transactions; presupuesto por usuario y mes.", "BIGINT positivo, moneda ISO, límites y checks de interés."],
        ["Automatización", "bank_ingestion_events 0 a 1 transactions; proposed_actions representa confirmación.", "event_id, vínculo de ingesta y hash de solicitud únicos."],
    ], [1.35, 3.4, 1.75])
    doc.add_page_break()

    heading(doc, "10 API, seguridad y privacidad")
    table(doc, ["Endpoint", "Autenticación", "Uso"], [
        ["GET /health", "Ninguna", "Liveness mínimo del backend."],
        ["GET /v1/agents/status", "Bearer API", "Estado de backend, LLM, correo, base y modelo."],
        ["POST /v1/assistant/messages", "Bearer API", "Mensaje, request_id y conversation_id opcional."],
        ["POST /v1/actions/{id}/confirm", "Bearer API", "Confirmación idempotente de propuesta."],
        ["GET /v1/workspace", "Bearer API", "Snapshot mensual autenticado."],
        ["GET PATCH /v1/profile y memorias", "Bearer API", "Perfil y memoria personal explícita."],
        ["POST /v1/webhooks/bank-events", "Webhook token", "Ingesta bancaria automatizada."],
    ], [2.0, 1.35, 3.2])
    table(doc, ["Riesgo", "Control implementado"], [
        ["Exposición del backend", "localhost más Tailscale Serve; sin Funnel ni puerto público."],
        ["Herramientas invocadas por el LLM", "Catálogo cerrado, esquemas Pydantic y reglas de dominio."],
        ["Cambios accidentales", "Propuestas explícitas, confirmación, recibos atómicos e idempotencia."],
        ["Secretos", "Variables y DPAPI fuera del repositorio; tokens separados y enmascarados."],
        ["Datos bancarios", "Persistencia limitada al hash y extracción estructurada necesaria."],
    ], [2.1, 4.45])

    heading(doc, "11 Validación y demostración")
    body(doc, "Las pruebas automatizadas validan dominio, cliente, backend, persistencia, propuestas, reintentos y estados móviles. La evidencia física necesaria para el taller se obtiene en un Android con dictado y Tailscale habilitados, una conexión por datos móviles y una automatización bancaria configurada.")
    table(doc, ["Nivel", "Cobertura", "Evidencia"], [
        ["Unitario", "Importes, fechas Bogotá, validadores y reglas de negocio.", "npm test."],
        ["Backend", "API, seguridad, herramientas, propuestas y persistencia.", "npm run test:backend."],
        ["Móvil", "Modo servidor, perfiles, memoria, retry, voz y propuestas.", "Jest y React Native Testing Library."],
        ["Dispositivo", "TTS, dictado, Tailscale remoto y webhook bancario.", "Video demostrativo de 3 a 5 minutos."],
    ], [1.1, 3.75, 1.7])
    heading(doc, "12 Limitaciones y evolución", 2)
    body(doc, "La versión está orientada a un propietario local estable y no declara soporte multiusuario. La configuración de automatización bancaria y correo es ambiental, por lo que requiere evidencia de dispositivo. Los adaptadores y contratos aíslan la UI, el proveedor LLM y la persistencia para permitir evolución sin reescribir los flujos principales.")
    heading(doc, "Referencias técnicas", 2)
    for item in (
        "Taller Segundo Corte Asistente Personal Inteligente Multi Agente con Control por Voz V2, secciones 1 a 7.",
        "Repositorio JARVIS: docs/ARCHITECTURE.md, docs/MULTI_AGENT_BACKEND.md, docs/JARVIS_VOICE.md y docs/WORKSHOP_REQUIREMENTS.md.",
        "Código fuente: backend/app, src/services/backend, src/services/storage y src/features/jarvis.",
    ): bullet(doc, item)
    doc.save(OUT)


if __name__ == "__main__":
    main()
